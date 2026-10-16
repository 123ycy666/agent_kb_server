"""Excel（.xlsx，新版 OOXML）解析器。

⭐ 逐行展开为「列名: 值」的规则见 `tabular_common` —— .xls 走的是同一套逻辑，
两边行为必须一致，否则会出现「同样的表格 .xls 能查、.xlsx 查不到」的怪问题。
"""

from __future__ import annotations

from pathlib import Path

from app.services.ingestion.parsers.base import ParsedDocument, ParserError
from app.services.ingestion.parsers.tabular_common import build_tabular_document


def parse_xlsx(path: Path) -> ParsedDocument:
    try:
        import openpyxl
    except ImportError as exc:
        raise ParserError(
            "缺少 Excel 解析依赖 openpyxl，请执行：pip install openpyxl"
        ) from exc

    try:
        # data_only=True：取公式的计算结果而非公式本身
        workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ParserError(
            f"Excel(.xlsx) 解析失败：{exc}。"
            "若该文件其实是旧版 .xls，请直接上传 .xls（同样支持），或用 Excel 另存为 .xlsx。"
        ) from exc

    try:
        sheets = [
            (name, workbook[name].iter_rows(values_only=True))
            for name in workbook.sheetnames
        ]
        return build_tabular_document(title=path.stem, sheets=sheets)
    finally:
        # ⚠️ read_only 模式下的 workbook 持有文件句柄，不关会在 Windows 上锁住文件，
        #    导致后续删除文档时 unlink 失败。
        workbook.close()
