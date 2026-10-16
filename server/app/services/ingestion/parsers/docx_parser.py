"""Word（.docx）解析器。

⭐ 关键点：保留标题层级。
python-docx 能识别 Heading 1~9 样式，据此构建 heading_path
（如「第三章 > 3.2 报销标准」），显著提升检索命中率。
只取纯文本会丢失这层信息 —— 这是最常见的解析失误。
"""

from __future__ import annotations

import re
from pathlib import Path

from app.services.ingestion.parsers.base import ParsedBlock, ParsedDocument, ParserError

_HEADING_STYLE = re.compile(r"^Heading\s*(\d)$", re.IGNORECASE)
# ⚠️ 仅「第X章 / 第X节」视为标题。**「第X条」是企业制度的条款正文本身**
#    （如「第三条 标准工作时间为…」），误判为标题会把后续条款挂到错误父节点下，
#    实测会污染 heading_path（前缀里混进上一句正文）。
_CN_HEADING = re.compile(r"^第[一二三四五六七八九十百零\d]+[章节]\s*")


def parse_docx(path: Path) -> ParsedDocument:
    try:
        import docx
    except ImportError as exc:
        raise ParserError(
            "缺少 Word 解析依赖 python-docx，请执行：pip install python-docx"
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise ParserError(
            f"Word 解析失败：{exc}。注意：仅支持 .docx 格式，"
            "旧版 .doc 请先另存为 .docx。"
        ) from exc

    doc = ParsedDocument(title=path.stem)
    heading_stack: dict[int, str] = {}

    def current_path() -> str | None:
        if not heading_stack:
            return None
        return " > ".join(heading_stack[k] for k in sorted(heading_stack))

    heading_count = 0
    table_count = 0

    for para in document.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue

        style_name = (para.style.name if para.style else "") or ""
        level = 0

        m = _HEADING_STYLE.match(style_name)
        if m:
            level = int(m.group(1))
        elif style_name.lower() in ("title",):
            level = 1
        elif _CN_HEADING.match(text) and len(text) <= 60:
            # 中文文档常用「第三章 xxx」但未套用 Heading 样式，做一次兜底识别
            level = 2

        if level:
            heading_stack = {k: v for k, v in heading_stack.items() if k < level}
            heading_stack[level] = text
            heading_count += 1
            doc.blocks.append(
                ParsedBlock(
                    text=text,
                    heading_path=" > ".join(
                        heading_stack[k] for k in sorted(heading_stack) if k < level
                    )
                    or None,
                    block_type="heading",
                    level=level,
                )
            )
            if doc.title == path.stem:
                doc.title = text
            continue

        doc.blocks.append(
            ParsedBlock(text=text, heading_path=current_path(), block_type="paragraph")
        )

    # 表格
    for table in document.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [(c.text or "").replace("\n", " ").strip() for c in row.cells]
            # 跳过整行空白
            if any(cells):
                rows.append(" | ".join(cells))
        if len(rows) >= 2:
            table_count += 1
            doc.blocks.append(
                ParsedBlock(
                    text="\n".join(rows),
                    heading_path=current_path(),
                    block_type="table",
                )
            )

    if doc.is_empty:
        raise ParserError(
            "该 Word 文档中未提取到任何文本内容（可能全为图片或文本框）"
        )

    doc.meta["heading_count"] = heading_count
    doc.meta["table_count"] = table_count
    doc.meta["paragraph_count"] = len(document.paragraphs)
    return doc
