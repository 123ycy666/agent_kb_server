"""公文审校回归（`proofread`）：规则必须全中，模型层只验接线。

跑法（仓库根）：
    py -X utf8 server/tests/agent/test_proofread.py

为什么不评"命中率"：规则层是确定性的，报错就是真错（下面逐条断言）；
模型层的召回/误报要用**真样本 + 真模型**评，见实施计划 V 判据，这里只验：
模型返回的条目能被合并、能回指原文、能进清单，且编造的条目标成"待确认"。
"""

from __future__ import annotations

import json
import os
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve()
ROOT = HERE.parents[3]
TMP = ROOT / ".dsh" / "tmp" / "proofread_test"
if TMP.exists():
    shutil.rmtree(TMP, ignore_errors=True)
TMP.mkdir(parents=True, exist_ok=True)
os.environ["DOC_FILES_DIR"] = str(TMP / "files")
sys.path.insert(0, str(ROOT / "server"))

from app.secretary import doc_files, pipeline_proofread     # noqa: E402

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(f"{name} {detail}".strip())
    print(("  PASS " if cond else "  FAIL ") + name + (f" | {detail}" if detail else ""))


def make_docx(path: Path) -> None:
    """故意塞进 6 类典型问题：错别字、英文标点、多余空格、百分号全角、
    单位混用、一级序号混用，外加一句病句（模型类）。"""
    import docx

    d = docx.Document()
    d.add_heading("关于加强配网运维管理的通知", level=1)
    d.add_paragraph("公司各单位,现将有关事项通知如下。")          # 英文逗号
    d.add_heading("一、总体要求", level=2)
    d.add_paragraph("各单位要结合实际制定实施方案,确保闭环率达到95％以上。")  # 英文逗号 + 全角％
    d.add_paragraph("要建立台账管理机制.明确责任分工。")             # 英文句点
    d.add_heading("二、资金安排", level=2)
    d.add_paragraph("本次安排资金120万元，另有机动车维修费用5000元。")  # 万元/元混用
    d.add_heading("第三章 责任分工", level=2)                      # 序号混用（第三章 vs 一、二、）
    d.add_paragraph("运检部负责统筹协调，各供电所负责具体实施。")
    d.add_paragraph("对于其它事项，由运检部另行通知。")               # 其它
    d.add_paragraph("截止目前，各项工作正在有序推进。")               # 截止目前
    d.add_paragraph("各单位  要按照要求报送材料。")                  # 多余空格
    d.save(str(path))


def fake_llm(messages, max_tokens=1200) -> str:
    if "公文审校" in messages[0]["content"]:
        return json.dumps({"issues": [
            {"type": "病句", "quote": "对于其它事项，由运检部另行通知",
             "suggestion": "其他事项由运检部另行通知", "reason": "「对于……，由……」句式啰嗦"},
            {"type": "表述不清", "quote": "这句话原文里没有",
             "suggestion": "（示例）", "reason": "用来验未核实标记"},
        ]}, ensure_ascii=False)
    return "{}"


def find(issues: list[dict], typ: str, quote_part: str = "") -> dict | None:
    for it in issues:
        if typ in it["type"] and (not quote_part or quote_part in (it.get("quote") or "")):
            return it
    return None


def main() -> int:
    src = TMP / "通知.docx"
    make_docx(src)
    f = doc_files.save(21, "通知.docx", src.read_bytes())

    pipeline_proofread.set_llm(fake_llm)
    res = pipeline_proofread.run(f["file_id"], 21, {}, lambda p, s: None, TMP / "out")
    issues = res["issues"]

    print("[1] 规则层（确定性，必须全中）")
    check("查出英文逗号", find(issues, "英文逗号") is not None)
    check("查出英文句点", find(issues, "英文句点") is not None)
    check("查出「其它」", (find(issues, "错别字", "其它") or {}).get("suggestion") == "其他")
    # 规则表里「截止目前」→「截至目前」、「截止到」→「截至」，两者都以「截至」开头
    check("查出「截止目前」→截至",
          ((find(issues, "错别字", "截止目前") or {}).get("suggestion") or "").startswith("截至"),
          str((find(issues, "错别字", "截止目前") or {}).get("suggestion")))
    check("查出全角百分号", find(issues, "数字符号") is not None)
    check("查出万元/元混用", find(issues, "单位混用") is not None)
    check("查出层次序号混用", find(issues, "层次序号混用") is not None)
    check("查出多余空格", find(issues, "多余空格") is not None)

    print("[2] 模型层（只验接线与自检）")
    check("模型报的病句进来了", find(issues, "病句") is not None)
    fabricated = find(issues, "表述不清")
    check("编造条目被标未核实", fabricated is not None and fabricated["verified"] is False,
          str(fabricated and fabricated.get("verified")))
    real = find(issues, "病句")
    check("真条目已核实", bool(real and real.get("verified")), str(real and real.get("verified")))

    print("[3] 去重与统计")
    keys = [(it["type"], (it.get("quote") or "")[:12]) for it in issues]
    check("没有完全重复的条目", len(keys) == len(set(keys)), f"{len(keys)} 条")
    check("分类统计存在", bool(res["stats"]["by_type"]), str(res["stats"]["by_type"]))
    check("总数一致", res["stats"]["total"] == len(issues))

    print("[4] 产出")
    md = res["answer_markdown"]
    check("markdown 有表格", "| # | 类型 |" in md)
    check("markdown 有统计", "分类统计" in md)
    check("产出 docx", len(res["outputs"]) == 1, str(res["outputs"]))
    out = TMP / "out" / res["outputs"][0]
    check("docx 存在且非空", out.exists() and out.stat().st_size > 5000,
          f"{out.stat().st_size if out.exists() else 0}B")
    import docx as _docx
    d = _docx.Document(str(out))
    check("docx 有清单表", len(d.tables) >= 1)
    if d.tables:
        head = [c.text for c in d.tables[0].rows[0].cells]
        check("表头正确", head[:3] == ["序号", "类型", "原文（节选）"], str(head))
        check("表里有数据", len(d.tables[0].rows) >= 5, f"rows={len(d.tables[0].rows)}")

    print()
    print(f"结果：{len(PASS)} 通过 / {len(FAIL)} 失败")
    if FAIL:
        print("失败项：")
        for x in FAIL:
            print("  -", x)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
