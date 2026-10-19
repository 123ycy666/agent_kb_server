"""文档任务接口（`/api/docs/jobs`）回归：建任务 → 跑完 → 下载成果。

跑法（仓库根）：
    py -X utf8 server/tests/agent/test_doc_jobs.py

本机沙箱起不了 PG，所以 `_resolve_kb_scope` 也打桩（与本轮要验的任务链路无关）。
模型用假模型（`pipeline_digest.set_llm`），要验的是**任务状态机 + 成果下载 + 归属校验**。
"""

from __future__ import annotations

import json
import os
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve()
ROOT = HERE.parents[3]
TMP = ROOT / ".dsh" / "tmp" / "jobs_test"
if TMP.exists():
    shutil.rmtree(TMP, ignore_errors=True)
TMP.mkdir(parents=True, exist_ok=True)
os.environ["DOC_FILES_DIR"] = str(TMP / "files")
os.environ["DOC_JOBS_DIR"] = str(TMP / "jobs")
sys.path.insert(0, str(ROOT / "server"))

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(f"{name} {detail}".strip())
    print(("  PASS " if cond else "  FAIL ") + name + (f" | {detail}" if detail else ""))


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("关于加强配网运维管理的通知", level=1)
    d.add_heading("一、时间要求", level=2)
    d.add_paragraph("请于2026年10月31日前将自查报告报送运检部。")
    d.add_heading("二、责任分工", level=2)
    d.add_paragraph("运检部负责统筹，各供电所负责实施。")
    d.save(str(path))


def fake_llm(messages, max_tokens=2000) -> str:
    if "公文分析助手" in messages[0]["content"]:
        return json.dumps({"items": [
            {"matter": "报送自查报告", "unit": "各供电所", "deadline": "2026年10月31日前",
             "materials": ["自查报告"], "quote": "请于2026年10月31日前将自查报告报送运检部"}
        ], "risks": []}, ensure_ascii=False)
    return "## 结论\n需在10月31日前报送自查报告。\n\n## 待办清单\n- [ ] 报送自查报告｜各供电所｜10月31日前｜自查报告"


def main() -> int:
    from fastapi.testclient import TestClient

    from app.core.deps import CurrentActor, get_current_actor
    from app.main import app
    from app.secretary import doc_files, pipeline_digest
    import app.api.v1.agent as agent_api

    pipeline_digest.set_llm(fake_llm)

    src = TMP / "通知.docx"
    make_docx(src)
    f = doc_files.save(11, "通知.docx", src.read_bytes())

    actors = {"me": CurrentActor(user_id=11, username="u11", display_name="甲"),
              "other": CurrentActor(user_id=12, username="u12", display_name="乙")}
    current = {"a": actors["me"]}

    app.dependency_overrides[get_current_actor] = lambda: current["a"]
    orig_scope = agent_api._resolve_kb_scope
    agent_api._resolve_kb_scope = lambda ids: ([], False)
    client = TestClient(app)
    try:
        print("[1] 建任务")
        r = client.post("/api/docs/jobs", json={"task": "digest", "file_ids": [f["file_id"]]})
        check("返回 202", r.status_code == 202, str(r.status_code))
        job = r.json()
        jid = job["job_id"]
        check("返回 job_id", jid.startswith("j_"), jid)

        print("[2] 轮询到完成")
        d = client.get(f"/api/docs/jobs/{jid}").json()
        check("状态 done", d["status"] == "done", f"{d['status']} / {d.get('stage')} / {d.get('error')}")
        check("进度 100", d["progress"] == 100, str(d["progress"]))
        check("有结果 markdown", "## 结论" in (d.get("result") or {}).get("answer_markdown", ""))
        check("有两个成果", len(d["outputs"]) == 2, str([o["name"] for o in d["outputs"]]))

        print("[3] 下载成果")
        name = d["outputs"][0]["name"]
        dr = client.get(f"/api/docs/jobs/{jid}/download", params={"name": name})
        check("下载 200", dr.status_code == 200, str(dr.status_code))
        check("文件非空", len(dr.content) > 5000, f"{len(dr.content)}B")

        print("[4] 列表与归属")
        check("列表里有它", client.get("/api/docs/jobs").json()["total"] >= 1)
        current["a"] = actors["other"]
        check("别人看不到", client.get(f"/api/docs/jobs/{jid}").status_code == 404,
              str(client.get(f"/api/docs/jobs/{jid}").status_code))
        check("别人下不了", client.get(f"/api/docs/jobs/{jid}/download",
                                     params={"name": name}).status_code == 404)
        current["a"] = actors["me"]

        print("[5] 参数与边界")
        bad = client.post("/api/docs/jobs", json={"task": "no_such_task", "file_ids": [f["file_id"]]})
        check("未支持的类型被拒", bad.status_code == 400, str(bad.status_code))
        empty = client.post("/api/docs/jobs", json={"task": "digest", "file_ids": []})
        check("空文件被拒", empty.status_code == 400, str(empty.status_code))
        check("不存在的任务 404", client.get("/api/docs/jobs/j_000000000000").status_code == 404)
        traversal = client.get(f"/api/docs/jobs/{jid}/download", params={"name": "../job.json"})
        check("路径穿越被挡", traversal.status_code == 404, str(traversal.status_code))
    finally:
        app.dependency_overrides.clear()
        agent_api._resolve_kb_scope = orig_scope

    print()
    print(f"结果：{len(PASS)} 通过 / {len(FAIL)} 失败")
    if FAIL:
        print("失败项：")
        for x in FAIL:
            print("  -", x)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
