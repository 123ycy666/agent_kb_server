# -*- coding: utf-8 -*-
"""F1-4 / F2-6 / F3-1 验收：写公文 Word、文件比对、Excel 直接问、纪要拟办。

这几项的共同点：**不依赖模型、不依赖新库**，本地就能跑到底。
（Excel 与拟办各自的细项见 test_excel_query.py / test_digest_to_task.py）

钉死五条
--------
1. 写公文：一份 docx → 按公文版式重排 → 走毕昇 `inspect_docx` 自检，结论"通过"。
2. 比对：两份材料 → 段落级差异 → 产出 md + docx；差异条数与人工构造的一致。
3. 比对无差异时：明确写"没有段落级差异"，不编。
4. 任务白名单：`deck` / `doc_write` / `compare` 都在 `/api/docs/jobs` 的可选任务里。
5. 拟办：digest 事项 → 时限不猜、缺单位标待指派、未核实标 need_review（委测在另一个文件）。

运行：py -X utf8 server/tests/agent/test_batch_features.py
"""
from __future__ import annotations

import shutil
import sys
import types
from pathlib import Path

SERVER = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(SERVER))

from app.secretary import office_runner as R  # noqa: E402

_checks: list[tuple[str, bool, str]] = []


def chk(label: str, cond: bool, extra: str = "") -> None:
    _checks.append((label, bool(cond), extra))


BLOCKS_A = [
    {'text': '关于加强配网改造管理的通知', 'block_type': 'heading', 'level': 1},
    {'text': '一、总体要求。各单位要高度重视配网改造工作，确保按期完成。', 'block_type': 'paragraph'},
    {'text': '二、进度安排。12 月底前完成全部项目验收。', 'block_type': 'paragraph'},
]
BLOCKS_B = [
    {'text': '关于加强配网改造管理的通知', 'block_type': 'heading', 'level': 1},
    {'text': '一、总体要求。各单位要高度重视配网改造工作，确保按期完成。', 'block_type': 'paragraph'},
    {'text': '二、进度安排。11 月底前完成全部项目验收，逾期通报。', 'block_type': 'paragraph'},
    {'text': '三、责任分工。运检部牵头，营销部配合。', 'block_type': 'paragraph'},
]


def _stub_doc_files(mapping: dict):
    """用桩替换 doc_files（真模块拉 sqlalchemy，本机没装）；接口与真模块一致。"""
    fake = types.ModuleType('app.secretary.doc_files')

    def load(fid, uid=None):
        return {'file_id': fid, 'name': mapping[fid]['name'], 'user_id': uid or 1}

    def blocks(fid, uid=None, offset=0, limit=50):
        return {'blocks': mapping[fid]['blocks'], 'total': len(mapping[fid]['blocks'])}

    fake.load = load
    fake.blocks = blocks
    sys.modules['app.secretary.doc_files'] = fake
    return fake


def main() -> int:
    print("批量功能验收：写公文 / 比对 / Excel / 拟办")
    work = SERVER / '.tmp' / 'batch_features'
    if work.exists():
        shutil.rmtree(work, ignore_errors=True)
    work.mkdir(parents=True)

    _stub_doc_files({'f_a': {'name': '通知（旧）.docx', 'blocks': BLOCKS_A},
                     'f_b': {'name': '通知（新）.docx', 'blocks': BLOCKS_B}})

    # ── 1 写公文 Word（走毕昇版式 + inspect_docx 闸门）────────────
    from app.secretary import pipeline_doc_write as W
    out1 = work / 'doc'
    out1.mkdir()
    r1 = W.run('f_a', 1, {}, lambda p, s: None, out1)
    docx_path = out1 / r1['outputs'][0]
    chk("写公文：产出 .docx", docx_path.is_file() and docx_path.suffix == '.docx', str(docx_path))
    chk("写公文：自检通过", r1['self_check']['verdict'] == '通过',
        r1['self_check']['tail'][-140:])
    import docx
    d = docx.Document(str(docx_path))
    text = '\n'.join(p.text for p in d.paragraphs)
    chk("写公文：标题与正文都在", '配网改造' in text and '总体要求' in text, text[:60])
    chk("写公文：末行有复核声明", '人工复核' in text)
    chk("写公文：格式来自站点包（仿宋/16磅）",
        r1['fmt']['body_font_cn'].startswith('仿宋') and r1['fmt']['body_pt'] == 16,
        str(r1['fmt']))

    # ── 2 文件比对 ────────────────────────────────────────────────
    from app.secretary import pipeline_compare as C
    out2 = work / 'cmp'
    out2.mkdir()
    r2 = C.run('f_a', 1, {'other_file_id': 'f_b'}, lambda p, s: None, out2)
    md = (out2 / '差异清单.md').read_text(encoding='utf-8')
    chk("比对：产出 md + docx", len(r2['outputs']) == 2 and (out2 / '差异清单.docx').is_file(),
        str(r2['outputs']))
    chk("比对：识别出 1 删 2 增",
        r2['summary'] == {'add': 2, 'del': 1, 'total': 3}, str(r2['summary']))
    chk("比对：md 里能看到新增段落", '责任分工' in md)
    chk("比对：md 里能看到删除段落", '12 月底前' in md)

    same = C.diff_paras(['甲', '乙'], ['甲', '乙'])
    chk("比对：相同输入 → 0 差异", same == [], str(same))
    md_same = C.to_markdown('a', 'b', [])
    chk("比对：无差异时明确写没有差异", '没有段落级差异' in md_same)

    # ── 3 Excel 直接问（openpyxl → sqlite3）──────────────────────
    from app.secretary import tools_excel as X
    book = work / '台账.xlsx'
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = '台账'
    for row in [('台区', '明细金额', '上报合计'), ('环潭', 100, 100), ('环潭', 50, 100),
                ('文峰', 200, 200)]:
        ws.append(list(row))
    wb.save(str(book))
    res = X.query(book, 'SELECT 台区, SUM(明细金额) AS s, MAX(上报合计) AS r FROM t '
                        'GROUP BY 台区 HAVING s <> r')
    chk("Excel：查出不平的台区", res['rows'] == [['环潭', 150, 100]], str(res['rows']))
    try:
        X.query(book, 'DELETE FROM t')
        blocked = False
    except X.QueryError:
        blocked = True
    chk("Excel：写语句被拒", blocked)

    # ── 4 任务白名单 ─────────────────────────────────────────────
    import ast, io
    src = io.open(str(SERVER / 'app/api/v1/doc_jobs.py'), encoding='utf-8').read()
    tree = ast.parse(src)
    tasks = set()
    for node in tree.body:
        if isinstance(node, ast.Assign) and getattr(node.targets[0], 'id', '') == 'TASKS':
            tasks = {k.value for k in node.value.keys}
    chk("白名单含 deck / doc_write / compare",
        {'deck', 'doc_write', 'compare'} <= tasks, str(sorted(tasks)))

    bad = [c for c in _checks if not c[1]]
    for label, ok, extra in _checks:
        print(("  [OK]   " if ok else "  [FAIL] ") + label + (("   ← " + extra) if not ok and extra else ""))
    print("\n%d 项检查，%d 项失败" % (len(_checks), len(bad)))
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
