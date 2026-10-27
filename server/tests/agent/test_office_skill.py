# -*- coding: utf-8 -*-
"""内置验收：agent 能不能真的跑技能包（bisheng-pptx）的脚本。

钉死的五条
----------
1. 模型写的构建脚本能在**我们的任务目录**里跑通，产物只出现在 `out/`。
2. 跑完能调**技能包自带的 `inspect_deck.py`** 做交付前体检，输出里有可判定的结论行。
3. 白名单生效：`import socket`（网络）、未登记模块、越界路径，全部拒。
4. 只允许跑 `skills/office/` 下的脚本（不许拿执行器跑任意文件）。
5. 超时能杀（这里用 1 秒超时跑一段 sleep 验证，不真等）。

⚠️ 这里用的构建脚本是**测试夹具**（按技能包 cookbook 的写法手写一条），
用来证明"执行器 + 技能脚本 + 自检闸门"这条链通了；真实运行时这段脚本由模型读
`SKILL.md` 后自己写。

运行：py -X utf8 server/tests/agent/test_office_skill.py
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))   # server/

from app.secretary import office_runner as R  # noqa: E402

BUILD_CODE = '''\
from pptx import Presentation
from pptx.util import Inches
import pptx_helpers as H

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

s1 = prs.slides.add_slide(prs.slide_layouts[0])
H.fill_text(s1.shapes.title, ['配网改造工作汇报'])
H.fill_text(s1.placeholders[1], ['2026 年 10 月', '小智 · 自动生成'])

s2 = prs.slides.add_slide(prs.slide_layouts[1])
H.fill_text(s2.shapes.title, ['一、总体情况'])
H.fill_text(s2.placeholders[1], ['本年度下达改造项目 12 个', '已完成 7 个，在建 5 个'])
prs.save('out/演示.pptx')
print('built', 'out/演示.pptx')
'''

_checks: list[tuple[str, bool, str]] = []


def chk(label: str, cond: bool, extra: str = "") -> None:
    _checks.append((label, bool(cond), extra))


def main() -> int:
    print("内置验收：技能包执行器 + 自检闸门")
    base = Path(__file__).resolve().parents[2] / '.tmp' / 'office_demo'
    if base.exists():
        shutil.rmtree(base, ignore_errors=True)
    wd = R.new_job_dir(base)
    chk("1 任务目录三层建好",
        all((wd / d).is_dir() for d in ('in', 'scratch', 'out')))

    # ── 1 跑构建脚本 ────────────────────────────────────
    res = R.run_code(BUILD_CODE, wd, timeout=120)
    chk("1b 构建脚本 exit=0", res['ok'], 'exit=%s stderr=%s' % (res['exit'], (res['stderr'] or '')[-300:]))
    outs = R.outputs(wd)
    chk("1c out/ 里只有一个交付物", len(outs) == 1 and outs[0]['name'] == '演示.pptx',
        str(outs))
    chk("1d 产物不是空文件", bool(outs) and outs[0]['bytes'] > 5000,
        str(outs[0]['bytes']) if outs else 'no file')

    # ── 2 技能包自带自检 ────────────────────────────────
    insp = R.run_script(R.PPTX_SCRIPTS / 'inspect_deck.py', ['out/演示.pptx'], wd, timeout=180)
    chk("2 inspect_deck 跑通（exit=0）", insp['ok'],
        'exit=%s %s' % (insp['exit'], (insp['stderr'] or '')[-300:]))
    tail = (insp['stdout'] or '')
    chk("2b 有逐页内容 dump", '[第2页]' in tail, tail[:120])
    chk("2c 有可判定的结论行", ('结论: 通过' in tail) or ('结论: 不通过' in tail), tail[-160:])
    chk("2d 本次样本判定为通过", '结论: 通过' in tail, tail[-160:])

    # ── 3 白名单 ────────────────────────────────────────
    for label, bad in [
        ('禁网络 import', 'import socket\nprint(1)'),
        ('禁未登记模块', 'import numpy\nprint(1)'),
        ('禁越界路径', "open('../外面.txt','w').write('x')"),
        ('空脚本', '   '),
    ]:
        try:
            R.run_code(bad, wd, timeout=30)
            ok = False
            why = '没被拒'
        except R.RunnerError as exc:
            ok = True
            why = str(exc)[:40]
        chk("3 拒绝：%s" % label, ok, why)

    # ── 4 只许跑技能包脚本 ──────────────────────────────
    try:
        R.run_script(Path(__file__), [], wd)
        ok = False
    except R.RunnerError:
        ok = True
    chk("4 拒绝跑 skills/office 之外的脚本", ok)

    # ── 5 超时 ──────────────────────────────────────────
    to = R.run_code('import time\ntime.sleep(5)', wd, timeout=1)
    chk("5 超时被杀（error=TIMEOUT）", to['error'] == 'TIMEOUT', str(to['error']))

    bad_list = [c for c in _checks if not c[1]]
    for label, ok, extra in _checks:
        print(("  [OK]   " if ok else "  [FAIL] ") + label + (("   ← " + extra) if not ok and extra else ""))
    print("\n%d 项检查，%d 项失败" % (len(_checks), len(bad_list)))
    return 1 if bad_list else 0


if __name__ == '__main__':
    sys.exit(main())
