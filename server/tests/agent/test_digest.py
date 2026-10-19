"""速读流水线回归（假模型，只验：切块→抽字段→归并→自检→出 docx/xlsx）。

跑法（仓库根）：
    py -X utf8 server/tests/agent/test_digest.py

为什么用假模型：这一层要验的是**流程与自检**（能不能切对块、出处能不能回指、
成果文件能不能打开），模型写得好不好不在这里判 —— 那要真样本 + 真模型，
见 docs/plans/2026-10-19-办公文档能力-实施计划.md 的 V2/V3。
"""

from __future__ import annotations

import os
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve()
ROOT = HERE.parents[3]
TMP = ROOT / ".dsh" / "tmp" / "digest_test"
if TMP.exists():
    shutil.rmtree(TMP, ignore_errors=True)
TMP.mkdir(parents=True, exist_ok=True)
os.environ["DOC_FILES_DIR"] = str(TMP / "files")
os.environ["DOC_JOBS_DIR"] = str(TMP / "jobs")
sys.path.insert(0, str(ROOT / "server"))

from app.secretary import doc_files, pipeline_digest       # noqa: E402

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(f"{name} {detail}".strip())
    print(("  PASS " if cond else "  FAIL ") + name + (f" | {detail}" if detail else ""))


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("关于开展2026年配网缺陷治理专项工作的通知", level=1)
    d.add_paragraph("公司各单位：现将有关事项通知如下。")

    d.add_heading("一、总体要求", level=2)
    d.add_paragraph("各单位要建立台账管理机制，确保缺陷闭环率达到95%以上，做到闭环管控。")

    d.add_heading("二、时间安排", level=2)
    d.add_paragraph("请于2026年11月30日前将治理台账报送运检部。")
    d.add_paragraph("每月5日前上报当月进展，重大缺陷24小时内报告。")

    d.add_heading("三、资金安排", level=2)
    d.add_paragraph("本次专项治理安排资金120万元，物资费用不超过60%。")

    d.add_heading("四、责任分工", level=2)
    d.add_paragraph("运检部负责统筹协调和进度考核，各供电所为实施主体，所长为第一责任人。")

    d.add_heading("五、考核与问责", level=2)
    d.add_paragraph("对逾期未报或台账不实的单位，将在月度考核中扣分并通报批评。")
    d.save(str(path))


FAKE_ITEMS = [
    {"matter": "建立台账管理机制", "unit": "各供电所", "deadline": "",
     "materials": [], "quote": "各单位要建立台账管理机制"},
    {"matter": "报送治理台账", "unit": "各供电所", "deadline": "2026年11月30日前",
     "materials": ["治理台账"], "quote": "请于2026年11月30日前将治理台账报送运检部"},
    {"matter": "上报当月进展", "unit": "各供电所", "deadline": "每月5日前",
     "materials": ["进展材料"], "quote": "每月5日前上报当月进展"},
    {"matter": "安排专项治理资金", "unit": "财务部", "deadline": "",
     "materials": [], "quote": "本次专项治理安排资金120万元"},
    {"matter": "落实责任分工", "unit": "运检部", "deadline": "",
     "materials": [], "quote": "运检部负责统筹协调和进度考核"},
    # 故意给一条原文里找不到的：验"未核实"标记，而不是被丢掉
    {"matter": "编制年度总结报告", "unit": "运检部", "deadline": "12月20日前",
     "materials": ["年度总结"], "quote": "这句话在原文里不存在"},
]


def fake_llm(messages, max_tokens=2000) -> str:
    sys_prompt = messages[0]["content"]
    if "公文分析助手" in sys_prompt:
        # 每块都返回全部候选，模拟"同一件事在多块里重复出现"，顺带验归并去重
        import json
        return json.dumps({"items": FAKE_ITEMS,
                           "risks": ["逾期未报将在月度考核中扣分并通报批评"]}, ensure_ascii=False)
    return ("## 结论\n这是配网缺陷治理专项通知，要求各单位建台账、按期限报送，资金 120 万元。\n\n"
            "## 一页纸要点\n- 闭环率 95% 以上\n- 11月30日前报台账\n\n"
            "## 待办清单\n- [ ] 报送治理台账｜各供电所｜2026年11月30日前｜治理台账\n\n"
            "## 风险与易漏\n- 逾期未报扣分并通报批评")


def main() -> int:
    src = TMP / "通知.docx"
    make_docx(src)
    f = doc_files.save(7, "通知.docx", src.read_bytes())

    pipeline_digest.set_llm(fake_llm)
    stages: list[str] = []
    out_dir = TMP / "out"
    res = pipeline_digest.run(f["file_id"], 7, {"want_docx": True, "want_xlsx": True},
                              lambda pct, stage: stages.append(f"{pct}:{stage}"), out_dir)

    print("[1] 抽取与归并")
    check("抽到事项", len(res["items"]) >= 4, f"items={len(res['items'])}")
    merged = [it for it in res["items"] if it["matter"] == "报送治理台账"]
    check("重复事项已归并", len(merged) == 1, f"命中 {len(merged)} 条")
    check("时限保留", merged and merged[0]["deadline"] == "2026年11月30日前")

    print("[2] 自检（出处回指）")
    verified = [it for it in res["items"] if it["verified"]]
    unverified = [it for it in res["items"] if not it["verified"]]
    check("真实原文的事项已核实", len(verified) >= 4, f"verified={len(verified)}")
    check("编造的事项标未核实", any(it["matter"] == "编制年度总结报告" for it in unverified),
          str([it["matter"] for it in unverified]))
    check("核实项带章节", all(it.get("section") for it in verified))
    check("warnings 提示未核实", any("未核实" in w for w in res["warnings"]), str(res["warnings"]))

    print("[3] 一页纸与清单")
    md = res["answer_markdown"]
    check("含结论段", "## 结论" in md)
    check("含待办清单段", "## 待办清单" in md)
    check("清单里带时限", "2026年11月30日前" in md)
    check("进度有推进记录", stages and stages[-1].endswith("完成"), str(stages[-1:]))

    print("[4] 成果文件")
    check("产出两个文件", len(res["outputs"]) == 2, str(res["outputs"]))
    docx_path = out_dir / res["outputs"][0]
    xlsx_path = out_dir / res["outputs"][1]
    check("docx 存在且非空", docx_path.exists() and docx_path.stat().st_size > 5000,
          f"{docx_path.stat().st_size if docx_path.exists() else 0}B")

    import docx as _docx
    d = _docx.Document(str(docx_path))
    check("docx 有表格", len(d.tables) >= 1, f"tables={len(d.tables)}")
    if d.tables:
        head = [c.text for c in d.tables[0].rows[0].cells]
        check("表头正确", head[:3] == ["事项", "责任单位", "时限"], str(head))
        check("表格有数据行", len(d.tables[0].rows) >= 3, f"rows={len(d.tables[0].rows)}")

    from openpyxl import load_workbook
    wb = load_workbook(str(xlsx_path))
    ws = wb.worksheets[0]
    check("xlsx 有表头", ws.cell(1, 1).value == "序号", str(ws.cell(1, 1).value))
    check("xlsx 冻结首行", ws.freeze_panes == "A2", str(ws.freeze_panes))
    check("xlsx 有数据", ws.max_row >= 3, f"max_row={ws.max_row}")

    print()
    print(f"结果：{len(PASS)} 通过 / {len(FAIL)} 失败")
    if FAIL:
        print("失败项：")
        for x in FAIL:
            print("  -", x)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
