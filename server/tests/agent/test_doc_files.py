"""doc_files 底座回归（不连模型、不连库，纯文件层）。

跑法（仓库根）：
    py -X utf8 server/tests/agent/test_doc_files.py

覆盖：
1. 上传 docx → 解析出标题/段落/表格，chars/blocks 合理；
2. 同内容重复上传 → 复用同一 file_id（不重复落盘）；
3. 别人的 file_id → FILE_NOT_FOUND（归属校验，不信入参）；
4. 不支持格式 → UNSUPPORTED_FORMAT；
5. blocks 分页、plain_text 带章节路径；
6. pptx（stdlib 解析）能取到每页标题与正文。
"""

from __future__ import annotations

import os
import shutil
import sys
import zipfile
from pathlib import Path

# 用工作区内的临时目录，别污染 data/doc_files
HERE = Path(__file__).resolve()
ROOT = HERE.parents[3]                      # server/tests/agent/xxx.py → 仓库根
TMP = ROOT / ".dsh" / "tmp" / "doc_files_test"
if TMP.exists():
    shutil.rmtree(TMP, ignore_errors=True)
TMP.mkdir(parents=True, exist_ok=True)

os.environ["DOC_FILES_DIR"] = str(TMP)
sys.path.insert(0, str(ROOT / "server"))

from app.secretary import doc_files                     # noqa: E402

PASS, FAIL = [], []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(f"{name} {detail}".strip())
    print(("  PASS " if cond else "  FAIL ") + name + (f" | {detail}" if detail else ""))


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("关于开展2026年配网缺陷治理工作的通知", level=1)
    d.add_paragraph("公司各单位：")
    d.add_heading("一、总体要求", level=2)
    d.add_paragraph("各单位要高度重视，确保缺陷闭环率达到 95% 以上。")
    d.add_heading("二、时间安排", level=2)
    d.add_paragraph("请于2026年10月31日前将治理台账报送运检部。")
    d.add_paragraph("联系人：张三，电话 1234567。")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "责任单位"
    t.cell(0, 1).text = "完成时限"
    t.cell(1, 0).text = "城东供电所"
    t.cell(1, 1).text = "2026-10-31"
    d.save(str(path))


def make_pptx(path: Path) -> None:
    """手工拼一个最小 pptx（zip + 2 页 XML），验证 stdlib 解析器。"""
    head = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<p:cSld><p:spTree>'
    )
    tail = "</p:spTree></p:cSld></p:sld>"

    def body(title: str, lines: list[str]) -> str:
        paras = "".join(
            f"<a:p><a:r><a:t>{t}</a:t></a:r></a:p>" for t in [title, *lines]
        )
        return f"<p:sp><p:txBody><a:bodyPr/>{paras}</p:txBody></p:sp>"

    with zipfile.ZipFile(path, "w") as z:
        z.writestr(
            "ppt/slides/slide1.xml",
            head + body("2026年缺陷治理汇报", ["一、总体情况", "闭环率 96%"]) + tail,
        )
        z.writestr(
            "ppt/slides/slide2.xml",
            head + body("二、下一步", ["继续压降缺陷存量"]) + tail,
        )


def main() -> int:
    src = TMP / "通知.docx"
    make_docx(src)
    content = src.read_bytes()

    print("[1] 上传与解析")
    f = doc_files.save(101, "通知.docx", content, user_code="zhangsan")
    check("返回 file_id", f["file_id"].startswith("f_"))
    check("解析成功", bool(f["parsed"].get("ok")), str(f["parsed"].get("error", "")))
    check("字数 > 100", f["parsed"].get("chars", 0) > 100, f"chars={f['parsed'].get('chars')}")
    check("有标题块", f["parsed"].get("headings", 0) >= 3, f"headings={f['parsed'].get('headings')}")
    check("有表格", f["parsed"].get("tables", 0) >= 1, f"tables={f['parsed'].get('tables')}")
    check("preview 非空", bool(f["parsed"].get("preview")))

    print("[2] 去重")
    f2 = doc_files.save(101, "另一个名字.docx", content, user_code="zhangsan")
    check("同内容复用 file_id", f2["file_id"] == f["file_id"], f2["file_id"])
    n_dirs = len(list((TMP / "u101").iterdir()))
    check("磁盘只有一份", n_dirs == 1, f"dirs={n_dirs}")

    print("[3] 归属校验")
    try:
        doc_files.load(f["file_id"], 999)
        check("别人的文件不可读", False)
    except doc_files.DocFileError as exc:
        check("别人的文件不可读", exc.code == "FILE_NOT_FOUND", exc.code)

    print("[4] 不支持的格式")
    try:
        doc_files.save(101, "virus.exe", b"MZ\x90\x00")
        check("拒绝 .exe", False)
    except doc_files.DocFileError as exc:
        check("拒绝 .exe", exc.code == "UNSUPPORTED_FORMAT", exc.code)

    print("[5] blocks / text")
    b = doc_files.blocks(f["file_id"], 101, offset=0, limit=3)
    check("blocks 分页", len(b["blocks"]) == 3 and b["total"] > 3, f"total={b['total']}")
    check("带 heading_path",
          any(x.get("heading_path") for x in doc_files.blocks(f["file_id"], 101)["blocks"]))
    text = doc_files.plain_text(f["file_id"], 101)
    check("正文含时限原句", "2026年10月31日前" in text)
    check("正文含表格内容", "城东供电所" in text)

    print("[6] pptx（stdlib 解析）")
    ppt = TMP / "汇报.pptx"
    make_pptx(ppt)
    fp = doc_files.save(101, "汇报.pptx", ppt.read_bytes(), user_code="zhangsan")
    check("pptx 可解析", bool(fp["parsed"].get("ok")), str(fp["parsed"].get("error", "")))
    ptxt = doc_files.plain_text(fp["file_id"], 101)
    check("第1页标题", "2026年缺陷治理汇报" in ptxt)
    check("第2页正文", "继续压降缺陷存量" in ptxt)
    check("页码定位", "第2页" in ptxt)

    print()
    print(f"结果：{len(PASS)} 通过 / {len(FAIL)} 失败")
    if FAIL:
        print("失败项：")
        for x in FAIL:
            print("  -", x)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
